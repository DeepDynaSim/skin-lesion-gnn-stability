from pathlib import Path
import sys,json,re,zipfile,hashlib
R=Path(__file__).resolve().parent
sys.path.insert(0,str(R.parent/'python_packages'));sys.path.insert(0,str(R.parent))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from manuscript_results_second import build

def main():
    values,tables=build();text=(R/'manuscript_template.md.in').read_text(encoding='utf-8');author='Caglar Uyulan\nFaculty of Engineering and Architecture, Department of Mechanical Engineering, İzmir Katip Celebi University, İzmir, Türkiye\nCorrespondence: caglar.uyulan@ikcu.edu.tr'
    for k,v in values.items():text=text.replace('{{'+k+'}}',v)
    text=text.replace('[AUTHOR]',author);assert '{{' not in text
    text=text.replace('seven regional controls, a descriptor ablation','eight regional controls')
    text=re.sub(r'^(### \d) (\d)',r'\1.\2',text,flags=re.M)
    dest=R/'final';dest.mkdir(exist_ok=True);doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85);sec.footer_distance=Inches(.3)
    for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[n];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);pr=st.element.get_or_add_rPr();color=pr.find(qn('w:color'))
        fonts=pr.find(qn('w:rFonts'))
        if fonts is not None:
            for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+a),None)
        if color is not None:
            for a in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+a),None)
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for b in list(pp):
                if b.tag==qn('w:pBdr'):pp.remove(b)
        st.paragraph_format.widow_control=True
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(10)
    for n,sz in [('Heading 1',13),('Heading 2',11.5)]:
        st=doc.styles[n];st.font.size=Pt(sz);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].paragraph_format.space_after=Pt(7)
    doc.core_properties.title='Stable graph propagation does not ensure image robustness: finite-step analysis and component evaluation in dermoscopic classification';doc.core_properties.author='Caglar Uyulan';doc.core_properties.subject='Finite-step graph theory, controlled lesion classification and external evaluation';doc.core_properties.keywords='dermoscopy; graph neural network; finite-step stability; calibration; PH2; reproducibility'
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    transform=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'));eqcount=figcount=tablecount=inlinecount=0;source=[]
    def math_xml(tex,size=21):
        xml=etree.fromstring(convert(tex).encode());math=transform(xml).getroot();assert math.tag.endswith('oMath')
        for rr in math.xpath('.//m:r',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}):
            pr=rr.find(qn('w:rPr'))
            if pr is None:pr=OxmlElement('w:rPr');rr.insert(1 if rr.find(qn('m:rPr')) is not None else 0,pr)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(size));pr.append(sz)
        return math
    def p(value,style=None):
        nonlocal inlinecount
        para=doc.add_paragraph(style=style)
        for segment in re.split(r'(\$[^$]+\$)',value):
            if segment.startswith('$') and segment.endswith('$'):para._p.append(math_xml(segment[1:-1],22));inlinecount+=1
            else:para.add_run(segment)
        if style is None:para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        source.append(value+'\n');return para
    def table(caption,headers,rows,widths):
        nonlocal tablecount;tablecount+=1;cap=p(f'Table {tablecount}. '+caption,'Caption');cap.paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        assert abs(sum(widths)-6.6)<.25,(caption,widths)
        for i,col in enumerate(t.columns):col.width=Inches(widths[i])
        for i,(cell,val) in enumerate(zip(t.rows[0].cells,headers)):cell.text=val;cell.width=Inches(widths[i])
        rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
        for row in rows:
            cells=t.add_row().cells
            for i,(cell,val) in enumerate(zip(cells,row)):cell.text=str(val);cell.width=Inches(widths[i])
        for ri,row in enumerate(t.rows):
            tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'C8C8C8');borders.append(el)
                pr.append(borders);mar=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
                pr.append(mar)
                if ri==0:el=OxmlElement('w:shd');el.set(qn('w:fill'),'EEEEEE');pr.append(el)
                for para in cell.paragraphs:
                    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1.02;para.paragraph_format.space_after=Pt(1)
                    para.paragraph_format.keep_with_next=(ri==0 or (len(t.rows)<=13 and ri<len(t.rows)-1))
                    if headers[ci]=='Evidence / manuscript section':para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for run in para.runs:run.font.size=Pt(8.8 if len(headers)>5 else 9.2);run.font.bold=ri==0
        source.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,row)).replace('\n',' ')+' |' for row in rows)+'\n');p('')
    for para in text.split('\n\n'):
        para=para.strip()
        if not para:continue
        if para.startswith('$$'):
            eqcount+=1;tex=para[2:-2];math=math_xml(tex)
            pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(7);pp.paragraph_format.keep_together=True;pp._p.append(math);run=pp.add_run(f'    ({eqcount})');run.font.size=Pt(10);source.append(para+f' ({eqcount})\n')
        elif para.startswith('# '):p(para[2:],'Title')
        elif para.startswith('## '):
            title=para[3:]
            if title.startswith('Supplementary'):doc.add_page_break()
            doc.add_heading(title,level=1);source.append(para+'\n')
        elif para.startswith('### '):doc.add_heading(para[4:],level=2);source.append(para+'\n')
        elif para.startswith('[TABLE:'):
            key=para[7:-1];table(*tables[key])
        elif para.startswith('[FIG:'):
            match=re.fullmatch(r'\[FIG:([^|]+)\|(.*)\]',para);assert match,para;name,caption=match.groups();figcount+=1;pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True;run=pp.add_run();run.add_picture(str(R/'figures'/(name+'.png')),width=Inches(6.65));run._r.xpath('.//wp:docPr')[0].set('descr',caption);p(f'Figure {figcount}. '+caption,'Caption');source.append(f'![Figure {figcount}](../figures/{name}.png)\n')
        elif para=='[REFERENCES]':
            refs=json.loads((R/'bibliography/references.json').read_text(encoding='utf-8'))
            for ref in refs:
                pp=p(f"[{ref['number']}] {ref['text']}");pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(3.5)
                for run in pp.runs:run.font.size=Pt(9.5)
        elif para=='[REPORTING_MAP]':
            doc.add_heading('Supplementary reporting maps',level=2);p('These abbreviated author-prepared mappings use the official 2024 item numbers. Locations refer to manuscript sections. Limited, unavailable and non-applicable items are retained; the maps do not certify clinical compliance. Full provenance and source URLs are supplied in S1.')
            table(*tables['demographics']);maps=json.loads((R/'analysis/reporting_maps.json').read_text())
            for key,title in [('TRIPOD_AI','TRIPOD+AI 2024: all 52 checklist subitems'),('CLAIM_2024','CLAIM 2024: all 44 checklist items')]:table(title,['Item','Topic','Status','Evidence / manuscript section'],maps[key],[.5,1.0,1.0,4.1])
        else:
            pp=p(para)
            if para.startswith('Caglar Uyulan\n'):
                pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in pp.runs:run.font.size=Pt(10.5)
    path=dest/'Skin_Lesion_GNN_Final_Manuscript.docx';doc.save(path);(dest/'Skin_Lesion_GNN_Final_Manuscript.md').write_text('\n'.join(source),encoding='utf-8')
    with zipfile.ZipFile(path) as z:
        xml=z.read('word/document.xml');root=etree.fromstring(xml);ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};assert len(root.findall('.//m:oMath',ns))==eqcount+inlinecount;assert b'{{' not in xml;assert b'[TABLE:' not in xml
    (dest/'document_manifest.json').write_text(json.dumps(dict(equations=eqcount,inline_equations=inlinecount,native_math=eqcount+inlinecount,figures=figcount,tables=tablecount,author='Caglar Uyulan',docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest()),indent=2));print(path,eqcount,inlinecount,figcount,tablecount,flush=True)
if __name__=='__main__':main()
